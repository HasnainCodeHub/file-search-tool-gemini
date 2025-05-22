import chainlit as cl
from agents import Agent, Runner, AsyncOpenAI, set_default_openai_api, set_default_openai_client, set_tracing_disabled
from agents.tool import function_tool
import os
from dotenv import load_dotenv
import mimetypes
from pdfminer.high_level import extract_text
import openpyxl
import docx
from openai.types.responses import ResponseTextDeltaEvent

# Load environment variables
load_dotenv()

# Get and validate Gemini API key
if not os.getenv("GEMINI_API_KEY"):
    raise ValueError("GEMINI_API_KEY is not set. Please ensure it is defined in your .env file.")
gemini_api_key = os.getenv('GEMINI_API_KEY')

# Set up Gemini API client
set_tracing_disabled(False)
external_client = AsyncOpenAI(
    api_key=gemini_api_key,
    base_url="https://generativelanguage.googleapis.com/v1beta/openai/",
)
set_default_openai_client(external_client)
set_default_openai_api("chat_completions")
set_tracing_disabled(True)

# Define the file_reader tool
@function_tool
def file_reader(file_path: str) -> str:
    """
    Reads content from a given file path and returns its textual content.
    Supports PDF, Excel, Word, and plain text files.
    """
    try:
        mime_type, _ = mimetypes.guess_type(file_path)

        if file_path.endswith(".pdf"):
            text = extract_text(file_path)
            return text or "No readable text found in PDF."

        elif file_path.endswith(".xlsx") or "spreadsheet" in str(mime_type):
            wb = openpyxl.load_workbook(file_path)
            content = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    content.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
            return "\n".join(content) or "No readable text found in Excel."

        elif file_path.endswith(".docx") or "word" in str(mime_type):
            doc = docx.Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs]) or "No readable text found in Word document."

        elif file_path.endswith(".txt") or "text" in str(mime_type):
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()

        else:
            return f"Unsupported file format: {file_path}"

    except Exception as e:
        return f"An error occurred while reading the file: {str(e)}"

# Set up the agent
agent = Agent(
    name="ResearcherBot",
    instructions=(
        "You are a research assistant. "
        "Use 'file_reader' to extract content from files when the user mentions reading or analyzing a file. "
        "Always try to show the full content."
        "You Can Also Response the General Pupose Questions If needed"
    ),
    tools=[file_reader],
    model="gemini-2.0-flash",
)

# Chainlit: Handle chat start
@cl.on_chat_start
async def start():
    """
    Handles the start of a chat session, prompting for and processing a file upload.
    """
    await cl.Message(content="Welcome! Please upload a file to analyze or ask a question.").send()

    files = await cl.AskFileMessage(
        content="Upload a file (PDF, Excel, Word, or text)",
        accept=["application/pdf", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "text/plain"],
        max_size_mb=10
    ).send()

    if files:
        file = files[0]
        file_path = file.path
        cl.user_session.set("file_path", file_path)
        await cl.Message(content=f"File uploaded: {file.name}. You can now ask questions about it.").send()
    else:
        await cl.Message(content="No file uploaded. You can still ask general questions.").send()

# Chainlit: Handle incoming messages
@cl.on_message
async def handle_message(message: cl.Message):
    """
    Handles user queries by processing them with the agent and streaming the response.
    """
    query = message.content
    file_path = cl.user_session.get("file_path")
    if file_path:
        query = f"The file is located here: {file_path}. {query}"

    # Create an initial empty message for streaming
    msg = cl.Message(content="")
    await msg.send()

    # Run the agent with streaming
    result = Runner.run_streamed(agent, query)
    async for event in result.stream_events():
        if event.type == "raw_response_event" and isinstance(event.data, ResponseTextDeltaEvent):
            # Stream each delta to the same message
            await msg.stream_token(event.data.delta)

    # Finalize the message
    await msg.update()

if __name__ == "__main__":
    cl.run()